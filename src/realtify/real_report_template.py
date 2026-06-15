from __future__ import annotations

import argparse
import json
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from rich.console import Console

from realtify.paths import PROJECT_ROOT
from realtify.report_template import inspect_report_template
from realtify.word_tools import ParagraphLocation, iter_paragraph_locations, set_paragraph_text


@dataclass(frozen=True)
class RealTemplateResult:
    template_path: Path
    inventory_json: Path
    inventory_md: Path
    replacements_count: int
    manual_review_count: int


def create_real_report_template(source_path: Path, output_path: Path) -> RealTemplateResult:
    source = _resolve_path(source_path)
    output = _resolve_path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)

    document = Document(str(output))
    replacements: list[dict[str, str | int]] = []
    manual_review: list[dict[str, str | int]] = []

    for location in iter_paragraph_locations(document):
        if location.area == "body" and location.index == 21 and _paragraph_has_highlight(location):
            replacements.append(
                {
                    "area": location.area,
                    "paragraph_index": location.index,
                    "old_text": location.paragraph.text,
                    "placeholder": "address_line_1_upper",
                }
            )
            set_paragraph_text(location.paragraph, "{{address_line_1_upper}}")
            for run in location.paragraph.runs:
                run.font.highlight_color = None
            continue
        highlighted_index = 0
        for run in location.paragraph.runs:
            if run.font.highlight_color is None or not run.text.strip():
                continue
            highlighted_index += 1
            placeholder = _placeholder_for_highlight(location, run.text, highlighted_index)
            if placeholder:
                replacements.append(
                    {
                        "area": location.area,
                        "paragraph_index": location.index,
                        "old_text": run.text,
                        "placeholder": placeholder,
                    }
                )
                run.text = f"{{{{{placeholder}}}}}"
                run.font.highlight_color = None
            else:
                manual_review.append(
                    {
                        "area": location.area,
                        "paragraph_index": location.index,
                        "text": run.text,
                    }
                )

    document.save(str(output))

    inventory = inspect_report_template(output)
    payload = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "source_path": str(source),
        "template_path": str(output),
        "replacements": replacements,
        "manual_review": manual_review,
        "template_inventory": inventory,
    }
    inventory_json = output.with_suffix(".inventory.json")
    inventory_md = output.with_suffix(".inventory.md")
    inventory_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    inventory_md.write_text(_inventory_markdown(payload), encoding="utf-8")

    return RealTemplateResult(
        template_path=output,
        inventory_json=inventory_json,
        inventory_md=inventory_md,
        replacements_count=len(replacements),
        manual_review_count=len(manual_review),
    )


def _placeholder_for_highlight(location: ParagraphLocation, text: str, highlighted_index: int) -> str | None:
    normalized = " ".join(text.replace("\u00a0", " ").split())
    lowered = normalized.casefold()

    if _looks_like_market_value_conclusion(normalized):
        return "market_value_uah_conclusion"
    if normalized == "1 376,00":
        return "median_price_usd_m2_report"
    if normalized == "42,3544":
        return "nbu_rate"
    if normalized == "58 279,65":
        return "median_price_uah_m2"
    if normalized == "2 727 487,62":
        return "market_value_uah_raw"
    if normalized == "375":
        return "apartment_number"
    if normalized == "46,8":
        return "total_area_m2"
    if normalized == "46,8 кв.м":
        return "total_area_m2_with_unit"
    if normalized == "ОДНОКІМНАТНОЇ":
        return "rooms_text_upper"
    if normalized == "однокімнатної":
        return "rooms_text"
    if normalized == "46,8 кв.м.":
        return "total_area_m2_with_unit"
    if normalized == "№ 433258095 від 27.06.2025 р.":
        return "extract_reference"

    if location.area == "body" and location.index == 21:
        return "address_line_1_upper"
    if location.area == "body" and location.index == 22:
        return "address_building"
    if location.area == "body" and location.index in {63, 65} and "2025" in normalized:
        return "valuation_date_long" if location.index == 63 else "report_date_long"
    if location.area == "body_table" and location.index == 574 and "2025" in normalized:
        return "valuation_date_long"

    if _looks_like_full_object_description(lowered):
        return "object_valuation_description"
    if _looks_like_short_object_description(lowered):
        return "object_valuation_description_short"
    if _looks_like_address(lowered):
        return "address_full" if location.area == "body" else "address_short"
    if "58 279,65" in normalized:
        return "median_price_uah_m2"
    return None


def _paragraph_has_highlight(location: ParagraphLocation) -> bool:
    return any(run.font.highlight_color is not None and run.text.strip() for run in location.paragraph.runs)


def _looks_like_address(lowered: str) -> bool:
    return (
        "\u043a\u0438\u0457\u0432" in lowered
        and "\u0434\u043d\u0456\u043f\u0440\u043e\u0432" in lowered
        and "17-\u043a" in lowered
    )


def _looks_like_short_object_description(lowered: str) -> bool:
    return "\u043e\u0434\u043d\u043e\u043a\u0456\u043c\u043d\u0430\u0442" in lowered and "375" in lowered and "46,8" in lowered


def _looks_like_full_object_description(lowered: str) -> bool:
    return _looks_like_short_object_description(lowered) and "\u0430\u0434\u0440\u0435\u0441\u043e\u044e" in lowered


def _looks_like_market_value_conclusion(text: str) -> bool:
    return text.startswith("2 727 500,00") and "\u0433\u0440\u043d" in text


def _inventory_markdown(payload: dict) -> str:
    lines = [
        "# Real Report Template Inventory",
        "",
        f"Source: {payload['source_path']}",
        f"Template: {payload['template_path']}",
        f"Replacements: {len(payload['replacements'])}",
        f"Manual review highlights: {len(payload['manual_review'])}",
        "",
        "## Replacements",
        "",
    ]
    for item in payload["replacements"]:
        lines.append(
            f"- {item['area']} paragraph {item['paragraph_index']}: "
            f"`{item['old_text']}` -> `{{{{{item['placeholder']}}}}}`"
        )
    if payload["manual_review"]:
        lines.extend(["", "## Manual Review", ""])
        for item in payload["manual_review"]:
            lines.append(f"- {item['area']} paragraph {item['paragraph_index']}: {item['text']}")
    placeholders = payload["template_inventory"].get("placeholders", [])
    if placeholders:
        lines.extend(["", "## Placeholders In Template", ""])
        lines.extend(f"- `{{{{{name}}}}}`" for name in placeholders)
    return "\n".join(lines).strip() + "\n"


def _resolve_path(path: Path) -> Path:
    return path if path.is_absolute() else PROJECT_ROOT / path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Create the real valuation .docx template from a highlighted report sample.")
    parser.add_argument("--source", type=Path, required=True, help="Converted .docx report sample.")
    parser.add_argument("--out", type=Path, required=True, help="Output .docx template path.")
    args = parser.parse_args(argv)

    console = Console()
    try:
        result = create_real_report_template(args.source, args.out)
    except Exception as exc:
        console.print(f"[red]Real report template creation failed:[/red] {exc}")
        return 1

    console.print("[green]Real report template created[/green]")
    console.print(f"Template: {result.template_path}")
    console.print(f"Inventory JSON: {result.inventory_json}")
    console.print(f"Inventory MD: {result.inventory_md}")
    console.print(f"Replacements: {result.replacements_count}")
    console.print(f"Manual review highlights: {result.manual_review_count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
